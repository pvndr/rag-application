from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.application.errors import (
    DocumentUploadError,
    DocumentExtractionError,
    EmptyDocumentError,
    UnsupportedFileTypeError,
)
from app.domain.repositories import DocumentExtractor


class MultiFormatDocumentExtractor(DocumentExtractor):
    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
    SUPPORTED_CONTENT_TYPES = {
        "application/pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "text/plain",
    }

    def extract(self, file_path: Path, content_type: str) -> str:
        extension = file_path.suffix.lower()
        if extension not in self.SUPPORTED_EXTENSIONS:
            raise UnsupportedFileTypeError(
                "Unsupported file type. Upload a PDF, DOCX, or TXT file."
            )

        try:
            if extension == ".pdf":
                text = self._extract_pdf(file_path)
            elif extension == ".docx":
                text = self._extract_docx(file_path)
            else:
                text = self._extract_txt(file_path)
        except DocumentUploadError:
            raise
        except Exception as exc:
            raise DocumentExtractionError(
                "Could not extract text from the uploaded document."
            ) from exc

        clean_text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
        if not clean_text:
            raise EmptyDocumentError("No readable text was found in the uploaded document.")

        return clean_text

    def _extract_pdf(self, file_path: Path) -> str:
        reader = PdfReader(str(file_path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    def _extract_docx(self, file_path: Path) -> str:
        document = DocxDocument(str(file_path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_cells = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
        return "\n".join(paragraphs + table_cells)

    def _extract_txt(self, file_path: Path) -> str:
        try:
            return file_path.read_text(encoding="utf-8")
        except UnicodeDecodeError as exc:
            raise DocumentExtractionError(
                "TXT uploads must be UTF-8 encoded."
            ) from exc
